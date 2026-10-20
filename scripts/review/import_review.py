#!/usr/bin/env python3
"""
確認済みの Excel（export_review.py で作ったもの）から、修正を public/locales/<言語>.json に取り込む（開発用）。
「Verdict」が Needs change で「Corrected translation」に入力がある行、または修正欄に入力がある行を反映する。
その後 `npm run check:locales` で、{置き換え} などが壊れていないか確かめる。

使い方: python3 scripts/review/import_review.py <ファイル.xlsx> [--locales <ディレクトリ>] [--dry-run]
"""
import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook')
    parser.add_argument('--locales', default=str(ROOT / 'public' / 'locales'))
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    match = re.match(r'sivree-review-(.+)\.xlsx$', Path(args.workbook).name)
    if not match:
        raise SystemExit('ファイル名は sivree-review-<言語>.xlsx の形にしてください')
    code = match.group(1)
    path = Path(args.locales) / f'{code}.json'
    messages = json.loads(path.read_text(encoding='utf-8'))

    sheet = load_workbook(args.workbook)['Review']
    header = [cell.value for cell in sheet[1]]
    col = {name: index for index, name in enumerate(header)}
    verdicts = {'OK': 0, 'Needs change': 0, '(blank)': 0}
    changed, skipped = [], []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        identifier, verdict = row[col['ID']], (row[col['Verdict']] or '').strip()
        corrected = (row[col['Corrected translation']] or '')
        corrected = corrected if isinstance(corrected, str) else str(corrected)
        verdicts[verdict if verdict in verdicts else '(blank)'] += 1
        if not corrected.strip():
            if verdict == 'Needs change':
                skipped.append(f'#{row[col["#"]]} 「{row[col["English (original)"]][:40]}」: Needs change だが修正欄が空')
            continue
        key, _, form = identifier.rpartition('|')
        current = messages[key]
        # 先頭・末尾の空白は意味を持つことがある（文の続きに付く文言）ため、原文に合わせて残す
        original = row[4]  # 5 列目が現在の訳（列の名前に言語名が入るため、位置で指定する）
        text = corrected
        if isinstance(original, str):
            if original.startswith(' ') and not text.startswith(' '):
                text = ' ' + text.lstrip()
            if original.endswith(' ') and not text.endswith(' '):
                text = text.rstrip() + ' '
        previous = current if not form else current.get(form)
        if previous == text:
            continue
        if form:
            current[form] = text
        else:
            messages[key] = text
        changed.append(f'#{row[col["#"]]} {key}{"|" + form if form else ""}: {previous!r} → {text!r}')

    print(f'{code}: OK {verdicts["OK"]} / Needs change {verdicts["Needs change"]} / 未回答 {verdicts["(blank)"]}')
    for line in changed:
        print('  変更', line)
    for line in skipped:
        print('  注意', line)
    if args.dry_run:
        print(f'（dry-run）{len(changed)} 件の変更は書き込みません')
        return
    if changed:
        path.write_text(json.dumps(messages, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'{len(changed)} 件を {path} に反映しました。次に npm run check:locales を実行してください')


if __name__ == '__main__':
    main()
